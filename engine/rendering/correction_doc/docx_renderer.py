"""DOCX rendering helpers for correction documents."""

from __future__ import annotations

import io
from typing import Dict, Optional, Tuple

from ...spec_engine.models import ChoiceRationale, PackagedQuiz, RationalesEntry
from .shared import _build_rationale_index, _choice_letter, _items_to_render, _strip_html
from .styles import (
    CORRECT_MARK_HEX,
    CORRECT_ROW_BG_HEX,
    HEADER_BG_HEX,
    HEADER_FG_HEX,
    INCORRECT_MARK_HEX,
    INCORRECT_ROW_BG_HEX,
    QUESTION_HDR_HEX,
    _COL_RATIOS,
    _FALLBACK_RATIONALE,
    _PAGE_WIDTH_INCHES,
)


def _shade_cell(cell, fill_hex: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:shd")):
        tcPr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def _set_col_widths(table, col_widths_inches: Tuple[float, ...]) -> None:
    from docx.shared import Inches

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(col_widths_inches):
                cell.width = Inches(col_widths_inches[i])


def _render_docx_item(doc, q_number: int, item: Dict, entry: RationalesEntry) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    hdr = doc.add_paragraph()
    run = hdr.add_run(f"Q{q_number}: {_strip_html(item.get('prompt', ''))}")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(QUESTION_HDR_HEX)
    run.font.size = Pt(11)

    if not entry.choices and getattr(entry, "text", None):
        qtype = item.get("type", "")
        label = "Model response to copy:" if qtype in ("ESSAY", "FILEUPLOAD") else "Explanation:"
        p = doc.add_paragraph()
        lab = p.add_run(label + " ")
        lab.bold = True
        lab.font.size = Pt(10)
        body = p.add_run(_strip_html(entry.text))
        body.font.size = Pt(10)
        doc.add_paragraph()
        return

    col_widths = tuple(r * _PAGE_WIDTH_INCHES for r in _COL_RATIOS)

    if entry.choices:
        choices_by_id = {c.id: c for c in entry.choices}
        item_choices = item.get("choices", [])

        table = doc.add_table(rows=1 + len(item_choices), cols=4)
        table.style = "Table Grid"

        hdr_row = table.rows[0]
        for cell, label in zip(hdr_row.cells, ["", "Choice", "✓ / ✗", "Rationale"]):
            _shade_cell(cell, HEADER_BG_HEX)
            p = cell.paragraphs[0]
            run = p.add_run(label)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(HEADER_FG_HEX)
            run.font.size = Pt(10)

        for row_idx, (choice_dict, data_row) in enumerate(zip(item_choices, table.rows[1:])):
            letter = _choice_letter(choice_dict, row_idx)
            is_correct = bool(choice_dict.get("correct"))
            choice_text = choice_dict.get("text", "")
            cr: Optional[ChoiceRationale] = choices_by_id.get(letter)
            rationale_text = cr.rationale if cr else _FALLBACK_RATIONALE

            row_bg = CORRECT_ROW_BG_HEX if is_correct else INCORRECT_ROW_BG_HEX
            text_color = CORRECT_MARK_HEX if is_correct else INCORRECT_MARK_HEX

            cells = data_row.cells
            _shade_cell(cells[0], row_bg)
            p0 = cells[0].paragraphs[0]
            r0 = p0.add_run(letter)
            r0.bold = True
            r0.font.color.rgb = RGBColor.from_string(text_color)
            r0.font.size = Pt(10)

            _shade_cell(cells[1], row_bg)
            p1 = cells[1].paragraphs[0]
            r1 = p1.add_run(choice_text)
            r1.bold = is_correct
            r1.font.color.rgb = RGBColor.from_string(text_color)
            r1.font.size = Pt(10)

            _shade_cell(cells[2], row_bg)
            p2 = cells[2].paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            mark = "✓" if is_correct else "✗"
            mark_color = CORRECT_MARK_HEX if is_correct else INCORRECT_MARK_HEX
            r2 = p2.add_run(mark)
            r2.bold = True
            r2.font.color.rgb = RGBColor.from_string(mark_color)
            r2.font.size = Pt(10)

            _shade_cell(cells[3], row_bg)
            p3 = cells[3].paragraphs[0]
            r3 = p3.add_run(rationale_text)
            r3.italic = not is_correct
            r3.font.color.rgb = RGBColor.from_string(text_color)
            r3.font.size = Pt(10)

        _set_col_widths(table, col_widths)

    doc.add_paragraph()


def render_docx_document(quiz: PackagedQuiz) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    title_text = quiz.title or "Correction Document"
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(f"{title_text} — Correction Document")
    title_run.bold = True
    title_run.font.size = Pt(15)
    title_run.font.color.rgb = RGBColor.from_string(HEADER_BG_HEX)

    subtitle_para = doc.add_paragraph("Review each answer choice and its explanation.")
    subtitle_para.runs[0].font.size = Pt(10)

    rationale_index = _build_rationale_index(quiz.rationales)
    renderable = _items_to_render(quiz.items, rationale_index)

    if not renderable:
        doc.add_paragraph(
            "No per-choice rationales found. Check that the quiz was generated "
            "with the per-choice rationale format."
        )
    else:
        for q_number, item, entry in renderable:
            _render_docx_item(doc, q_number, item, entry)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
