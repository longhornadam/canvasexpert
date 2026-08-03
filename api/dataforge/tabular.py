#!/usr/bin/env python3
"""Minimal spreadsheet/CSV reader covering exactly what the parsers use.

This replaces pandas. The parsers only ever needed positional cell access, a
row and column count, blank tests, and one header-row mode, so carrying pandas
plus numpy for that was a heavy dependency in a tool that installs into a
teacher's own user account with no admin rights.

`Grid` deliberately mimics the small slice of the DataFrame API the parsers
already used (`.iloc[r, c]`, `.shape`, `.columns`, `len()`, `.iterrows()`) so
the parsing logic did not have to be rewritten alongside the dependency change.

Two pandas behaviors are reproduced on purpose, because the parsers depend on
them:

  - An empty cell reads as None, and an empty string is an empty cell.
  - Trailing all-blank rows and columns are trimmed, so a sheet's shape
    reflects its content rather than whatever range the writer happened to
    touch.

One pandas behavior is deliberately NOT reproduced: per-column dtype. pandas
promotes an integer column to float64 if a single cell is blank, which turned
local ID 696969 into "696969.0" and silently re-keyed the anonymizer. Values
here are typed per cell, so that failure mode does not exist. Leading zeros in
an ID are preserved as text for the same reason.
"""

import csv as _csv
import math
import re
from pathlib import Path
from typing import Any, Iterator, List, Optional, Sequence, Tuple

import openpyxl

_INT_RE = re.compile(r"[+-]?\d+$")
_NON_FINITE = {"nan", "inf", "-inf", "+inf", "infinity", "-infinity", "+infinity"}


def is_blank(value: Any) -> bool:
    """True for an empty cell. The stand-in for pandas.isna."""
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def not_blank(value: Any) -> bool:
    """The stand-in for pandas.notna."""
    return not is_blank(value)


def coerce(text: Any) -> Any:
    """Type a raw text cell the way a reader should.

    Numeric strings become int or float; everything else stays text. Values
    with a leading zero stay text, because those are identifiers rather than
    quantities and losing the zero would change who they refer to.
    """
    if text is None:
        return None
    if not isinstance(text, str):
        return text
    s = text.strip()
    if not s:
        return None
    if s.lower() in _NON_FINITE:
        return s
    if _INT_RE.match(s):
        digits = s.lstrip("+-")
        if len(digits) > 1 and digits.startswith("0"):
            return s
        try:
            return int(s)
        except ValueError:
            return s
    try:
        return float(s)
    except ValueError:
        return s


def _trim(rows: List[List[Any]]) -> List[List[Any]]:
    """Drop trailing all-blank rows and columns, as pandas does on read."""
    while rows and all(is_blank(v) for v in rows[-1]):
        rows.pop()
    if not rows:
        return rows
    width = max(len(r) for r in rows)
    for r in rows:
        r.extend([None] * (width - len(r)))
    while width > 0 and all(is_blank(r[width - 1]) for r in rows):
        for r in rows:
            del r[width - 1]
        width -= 1
    return rows


class _ILoc:
    """Positional accessor, so `grid.iloc[row, col]` reads like it always did."""

    __slots__ = ("_rows",)

    def __init__(self, rows: List[List[Any]]):
        self._rows = rows

    def __getitem__(self, key: Tuple[int, int]) -> Any:
        if not isinstance(key, tuple) or len(key) != 2:
            raise TypeError("Grid.iloc takes a (row, column) pair")
        r, c = key
        # Out of range raises, matching the DataFrame behavior the callers
        # were written against. Silently returning None here would turn a
        # layout mismatch into a plausible-looking empty result.
        row = self._rows[r]
        return row[c]


class Grid:
    """A rectangular block of cells with optional column names."""

    def __init__(self, rows: List[List[Any]], columns: Optional[Sequence[Any]] = None):
        self._rows = rows
        ncols = len(rows[0]) if rows else (len(columns) if columns else 0)
        self._columns: List[Any] = list(columns) if columns is not None else list(range(ncols))

    @property
    def iloc(self) -> _ILoc:
        return _ILoc(self._rows)

    @property
    def columns(self) -> List[Any]:
        return self._columns

    @property
    def shape(self) -> Tuple[int, int]:
        return (len(self._rows), len(self._columns))

    def __len__(self) -> int:
        return len(self._rows)

    def iterrows(self) -> Iterator[Tuple[int, dict]]:
        """(index, {column_name: value}) per row, like DataFrame.iterrows."""
        for i, row in enumerate(self._rows):
            yield i, {name: row[j] if j < len(row) else None
                      for j, name in enumerate(self._columns)}


def read_xlsx(path, header: Optional[int] = None, nrows: Optional[int] = None) -> Grid:
    """Read the first worksheet.

    header=None gives a positional grid; header=0 takes the first row as
    column names. data_only=True so a formula cell yields its cached value
    rather than the formula text.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        limit = None if nrows is None else (nrows if header is None else nrows + 1)
        rows: List[List[Any]] = []
        for raw in ws.iter_rows(values_only=True):
            rows.append([None if (isinstance(v, str) and not v.strip()) else v for v in raw])
            if limit is not None and len(rows) >= limit:
                break
    finally:
        wb.close()

    rows = _trim(rows)
    if header is None:
        return Grid(rows)
    if not rows:
        return Grid([], [])
    names = [("" if is_blank(v) else v) for v in rows[0]]
    return Grid(rows[1:], names)


def read_csv(path, header: Optional[int] = 0) -> Grid:
    """Read a CSV. header=0 takes the first row as column names."""
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        raw = [[coerce(cell) for cell in row] for row in _csv.reader(fh)]
    raw = _trim(raw)
    if header is None:
        return Grid(raw)
    if not raw:
        return Grid([], [])
    names = [("" if is_blank(v) else str(v)) for v in raw[0]]
    return Grid(raw[1:], names)


def read_table(path, header: Optional[int] = 0) -> Grid:
    """Dispatch on suffix: .csv through the CSV reader, anything else as xlsx."""
    p = Path(path)
    if p.suffix.lower() == ".csv":
        return read_csv(p, header=header)
    return read_xlsx(p, header=header)
