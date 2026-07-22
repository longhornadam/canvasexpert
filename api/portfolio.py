"""Render a student's New Quizzes writing into a compiled portfolio DOCX.

Phase 1: builds a per-student document from parsed student_analysis data alone — an
MC performance summary plus the constructed (written) responses as prompt -> response
blocks. Phase 2 will merge live Assignment text entries + uploaded files/photos into
the same document (see the plan). Reuses python-docx (already a dependency; same
pattern as api/student_packet.py).

PII: portfolio documents are real student work — callers must write them only to the
synced student-reports root / gitignored output, never into the repo.
"""
import os
import re
from datetime import datetime

from docx import Document
from docx.shared import Inches

from api.nq_report import constructed_responses, html_to_text
from api.webui import workspace

_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}


def _safe(name: str, max_len: int = 80) -> str:
    """Filesystem-safe stem (local copy so this module stays import-light)."""
    cleaned = re.sub(r'[^\w\- ]+', "", (name or "").strip())
    cleaned = re.sub(r'\s+', " ", cleaned)
    return (cleaned[:max_len] or "student").strip()


def render_student_docx(student: dict, dest: str, quiz_title: str = "New Quiz") -> str:
    """Write one student's portfolio section for a single quiz to `dest`. Returns dest."""
    doc = Document()
    doc.add_heading(student.get("name", "Student"), level=0)

    meta = []
    if student.get("section"):
        meta.append(student["section"])
    if student.get("submitted"):
        meta.append(student["submitted"][:10])
    meta.append(f"generated {datetime.now():%b %d, %Y}")
    doc.add_paragraph(" · ".join(meta))

    doc.add_heading(quiz_title, level=1)

    # Multiple-choice performance summary (context, not writing).
    score = student.get("overall_score")
    poss = student.get("points_possible")
    summary = doc.add_table(rows=1, cols=4)
    summary.style = "Light Grid Accent 1"
    hdr = ["Score", "Correct (MC)", "Incorrect (MC)", "No response"]
    for i, h in enumerate(hdr):
        summary.rows[0].cells[i].text = h
    cells = summary.add_row().cells
    cells[0].text = f"{score:g}/{poss:g}" if score is not None and poss is not None else ""
    cells[1].text = str(student.get("num_correct") if student.get("num_correct") is not None else "")
    cells[2].text = str(student.get("num_incorrect") if student.get("num_incorrect") is not None else "")
    cells[3].text = str(student.get("no_response") if student.get("no_response") is not None else "")

    # Constructed (written) responses — the portfolio core.
    written = constructed_responses(student)
    doc.add_heading("Written responses", level=2)
    if not written:
        doc.add_paragraph("No written responses on this quiz.")
    for it in written:
        doc.add_heading(it.get("prompt") or "(untitled prompt)", level=3)
        body = html_to_text(it.get("response", ""))
        if body:
            for para in body.split("\n\n"):
                doc.add_paragraph(para)
        else:
            doc.add_paragraph("(no response)")
        ep = it.get("earned_points")
        if ep is not None:
            poss = it.get("points_possible_est")
            label = (f"Points earned: {ep:g} ({poss:g} possible)"
                     if poss is not None else f"Points earned: {ep:g}")
            doc.add_paragraph(label).italic = True

    os.makedirs(workspace.extended_path(os.path.dirname(os.path.abspath(dest))), exist_ok=True)
    doc.save(workspace.extended_path(dest))
    return dest


def render_portfolio(parsed: dict, out_dir: str, quiz_title: str = "New Quiz") -> list:
    """Render one DOCX per student from a parsed student_analysis dict.
    Returns the list of written file paths. `out_dir` must be outside the repo."""
    written = []
    for student in parsed.get("students", []):
        stem = _safe(student.get("name", "student"))
        dest = os.path.join(out_dir, f"{stem} - {_safe(quiz_title)}.docx")
        written.append(render_student_docx(student, dest, quiz_title))
    return written


# --------------------------------------------------------------------------
# Merged portfolio: one chronological document per student across sources
# (New Quizzes constructed responses + Assignment text entries + uploads/photos).
# Entry shape (built by the caller / portfolio_service):
#   {date: "YYYY-MM-DD"|"", source: str, title: str, prompt: str,
#    response_text: str, attachments: [local_path,...], earned: float|None,
#    possible: float|None}
# --------------------------------------------------------------------------

def _attachment_to_doc(doc, path):
    """Embed an attachment into the doc: images inline, PDF/DOCX as extracted text,
    anything else (or on failure) as a reference line. Heavy deps are lazy-imported
    so this module stays importable without them."""
    name = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    if ext in _IMAGE_EXTS:
        try:
            doc.add_picture(path, width=Inches(5.5))
            return
        except Exception:
            doc.add_paragraph(f"[Image attachment: {name} — could not embed]")
            return
    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(path)
            text = "\n".join((pg.extract_text() or "") for pg in reader.pages).strip()
            if text:
                for para in text.split("\n\n"):
                    doc.add_paragraph(para)
            else:
                doc.add_paragraph(f"[Attached PDF: {name} ({len(reader.pages)} page(s)) "
                                  "— no extractable text (likely scanned/photographed)]")
        except Exception:
            doc.add_paragraph(f"[Attached PDF: {name}]")
        return
    if ext == ".docx":
        try:
            sub = Document(path)
            for p in sub.paragraphs:
                if p.text.strip():
                    doc.add_paragraph(p.text)
        except Exception:
            doc.add_paragraph(f"[Attached document: {name}]")
        return
    doc.add_paragraph(f"[Attached file: {name}]")


def nq_entries(parsed: dict, student: dict, quiz_title: str) -> list:
    """Build portfolio entries from one student's New Quizzes data:
    an MC performance line + one entry per constructed (written) response."""
    date = (student.get("submitted") or "")[:10]
    entries = []
    nc, ni = student.get("num_correct"), student.get("num_incorrect")
    if nc is not None or ni is not None:
        score, poss = student.get("overall_score"), student.get("points_possible")
        summary = f"{nc or 0} correct / {ni or 0} incorrect"
        if score is not None and poss is not None:
            summary += f" · score {score:g}/{poss:g}"
        entries.append({"date": date, "source": "New Quiz",
                        "title": f"{quiz_title} — multiple choice", "prompt": "",
                        "response_text": summary, "attachments": [],
                        "earned": None, "possible": None})
    for it in constructed_responses(student):
        entries.append({"date": date, "source": "New Quiz", "title": quiz_title,
                        "prompt": it.get("prompt", ""),
                        "response_text": html_to_text(it.get("response", "")),
                        "attachments": [],
                        "earned": it.get("earned_points"),
                        "possible": it.get("points_possible_est")})
    return entries


def sort_entries(entries: list) -> list:
    """Chronological (oldest first); undated entries sort to the end."""
    return sorted(entries, key=lambda e: (e.get("date") or "9999-99-99", e.get("source", "")))


def render_merged_docx(student_name: str, entries: list, dest: str,
                       meta: str = "") -> str:
    """Render one student's combined, chronological writing portfolio."""
    doc = Document()
    doc.add_heading(student_name, level=0)
    doc.add_paragraph(meta or f"Writing portfolio · generated {datetime.now():%b %d, %Y}")
    if not entries:
        doc.add_paragraph("No writing on record for the selected sources.")
    for e in sort_entries(entries):
        head = " — ".join(p for p in [e.get("date"), e.get("title")] if p) or e.get("source", "Work")
        doc.add_heading(head, level=2)
        sub = e.get("source", "")
        if sub:
            doc.add_paragraph(sub).italic = True
        if e.get("prompt"):
            doc.add_paragraph(f"Prompt: {e['prompt']}").italic = True
        body = e.get("response_text") or ""
        for para in body.split("\n\n"):
            if para.strip():
                doc.add_paragraph(para)
        for att in e.get("attachments") or []:
            if os.path.exists(att):
                _attachment_to_doc(doc, att)
        earned, poss = e.get("earned"), e.get("possible")
        if earned is not None:
            label = (f"Points earned: {earned:g} ({poss:g} possible)"
                     if poss is not None else f"Points earned: {earned:g}")
            doc.add_paragraph(label).italic = True
    os.makedirs(workspace.extended_path(os.path.dirname(os.path.abspath(dest))), exist_ok=True)
    doc.save(workspace.extended_path(dest))
    return dest
