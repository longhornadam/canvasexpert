"""Source-material helpers for PowerGrader AI scoring.

Teachers usually keep passages as PDFs, Word docs, slide decks, or pasted text.
This module extracts plain text from those files and builds the shared context
block that travels once with a PowerGrader LLM request.
"""
from __future__ import annotations

import io
import json
import os
import re
import zipfile
from html import unescape
from pathlib import Path
from xml.etree import ElementTree as ET

from . import workspace


SOURCE_FOLDER_NAME = "Source Materials"
MAX_EXTRACTED_CHARS = 750_000

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".xml", ".yaml", ".yml",
    ".html", ".htm", ".rtf",
}
SUPPORTED_EXTS = TEXT_EXTS | {".pdf", ".docx", ".pptx", ".xlsx", ".odt"}
UNSUPPORTED_LEGACY_EXTS = {".doc", ".ppt", ".xls", ".pages", ".key", ".numbers"}

RESPONSE_PRESETS = {
    "scr": {
        "label": "SCR - single paragraph",
        "response_words": 130,
        "output_tokens_per_student": 350,
    },
    "ecr": {
        "label": "ECR - 4-5 paragraphs",
        "response_words": 650,
        "output_tokens_per_student": 600,
    },
}


def source_folder() -> str | None:
    return workspace.folder(SOURCE_FOLDER_NAME)


def ensure_source_folder() -> str | None:
    folder = source_folder()
    if folder:
        os.makedirs(folder, exist_ok=True)
    return folder


def _estimate_tokens_for_text(text: str) -> int:
    return max(0, len(text or "") // 4)


def estimate_text_tokens(text: str) -> int:
    return _estimate_tokens_for_text(text)


def list_source_files() -> list[dict]:
    folder = ensure_source_folder()
    if not folder or not os.path.isdir(folder):
        return []
    out: list[dict] = []
    root = Path(folder).resolve()
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_EXTS | UNSUPPORTED_LEGACY_EXTS:
            continue
        rel = str(path.relative_to(root))
        out.append({
            "name": path.name,
            "relpath": rel,
            "path": str(path),
            "size": path.stat().st_size,
            "supported": path.suffix.lower() in SUPPORTED_EXTS,
        })
    return out


def _decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _strip_html(text: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
    text = re.sub(r"(?s)<br\s*/?>", "\n", text)
    text = re.sub(r"(?s)</p\s*>", "\n\n", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return _collapse_ws(unescape(text))


def _strip_rtf(text: str) -> str:
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+\d* ?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return _collapse_ws(text)


def _collapse_ws(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in (text or "").splitlines()]
    compact: list[str] = []
    blank = False
    for line in lines:
        if not line:
            if not blank:
                compact.append("")
            blank = True
        else:
            compact.append(line)
            blank = False
    return "\n".join(compact).strip()


def _truncate(text: str) -> tuple[str, bool]:
    if len(text or "") <= MAX_EXTRACTED_CHARS:
        return text or "", False
    return (text or "")[:MAX_EXTRACTED_CHARS], True


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception as e:
        raise ValueError(f"PDF extraction requires pypdf: {e}") from e

    reader = PdfReader(io.BytesIO(data))
    chunks = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            chunks.append(f"[Page {i}]\n{page_text.strip()}")
    return _collapse_ws("\n\n".join(chunks))


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except Exception as e:
        raise ValueError(f"DOCX extraction requires python-docx: {e}") from e

    doc = Document(io.BytesIO(data))
    chunks = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return _collapse_ws("\n".join(chunks))


def _xml_text_nodes(raw: bytes) -> list[str]:
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return []
    texts = []
    for elem in root.iter():
        if elem.text and elem.text.strip() and elem.tag.rsplit("}", 1)[-1] in {"t", "p", "span"}:
            texts.append(elem.text.strip())
    return texts


def _extract_pptx(data: bytes) -> str:
    chunks = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = sorted(
            n for n in zf.namelist()
            if n.startswith("ppt/slides/slide") and n.endswith(".xml")
        )
        for i, name in enumerate(names, start=1):
            texts = _xml_text_nodes(zf.read(name))
            if texts:
                chunks.append(f"[Slide {i}]\n" + "\n".join(texts))
    return _collapse_ws("\n\n".join(chunks))


def _extract_xlsx(data: bytes) -> str:
    rows = []
    shared: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if "xl/sharedStrings.xml" in zf.namelist():
            try:
                root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
                for si in root:
                    parts = [node.text.strip() for node in si.iter()
                             if node.text and node.text.strip()]
                    if parts:
                        shared.append(" ".join(parts))
            except ET.ParseError:
                shared = []
        sheets = sorted(
            n for n in zf.namelist()
            if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")
        )
        for sheet_idx, name in enumerate(sheets, start=1):
            values = []
            try:
                root = ET.fromstring(zf.read(name))
            except ET.ParseError:
                continue
            for c in root.iter():
                if c.tag.rsplit("}", 1)[-1] != "c":
                    continue
                cell_type = c.attrib.get("t", "")
                text_value = ""
                v = next((child for child in c if child.tag.rsplit("}", 1)[-1] == "v"), None)
                if v is not None and v.text:
                    if cell_type == "s":
                        try:
                            text_value = shared[int(v.text)]
                        except (ValueError, IndexError):
                            text_value = v.text
                    else:
                        text_value = v.text
                inline = [node.text.strip() for node in c.iter()
                          if node.text and node.text.strip()
                          and node.tag.rsplit("}", 1)[-1] == "t"]
                if inline:
                    text_value = " ".join(inline)
                if text_value:
                    values.append(text_value)
            if values:
                rows.append(f"[Sheet {sheet_idx}]\n" + "\n".join(values))
    return _collapse_ws("\n\n".join(rows))


def _extract_odt(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if "content.xml" not in zf.namelist():
            return ""
        texts = _xml_text_nodes(zf.read("content.xml"))
    return _collapse_ws("\n".join(texts))


def extract_text_from_bytes(filename: str, data: bytes) -> tuple[str, list[str]]:
    ext = Path(filename or "").suffix.lower()
    warnings: list[str] = []
    if ext in UNSUPPORTED_LEGACY_EXTS:
        raise ValueError(
            f"{ext} files are old binary formats. Save as PDF, DOCX, PPTX, or XLSX first."
        )
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"Unsupported source-material file type: {ext or '(none)'}")

    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".pptx":
        text = _extract_pptx(data)
    elif ext == ".xlsx":
        text = _extract_xlsx(data)
    elif ext == ".odt":
        text = _extract_odt(data)
    else:
        text = _decode_bytes(data)
        if ext in {".html", ".htm"}:
            text = _strip_html(text)
        elif ext == ".rtf":
            text = _strip_rtf(text)
        else:
            text = _collapse_ws(text)

    text, truncated = _truncate(text)
    if truncated:
        warnings.append(
            f"{filename} was truncated to {MAX_EXTRACTED_CHARS:,} characters for safety."
        )
    if not text.strip():
        raise ValueError(f"No readable text could be extracted from {filename}.")
    return text, warnings


def _resolve_folder_file(relpath: str) -> Path:
    folder = ensure_source_folder()
    if not folder:
        raise ValueError("No Source Materials folder is available.")
    root = Path(folder).resolve()
    candidate = (root / relpath).resolve()
    if root not in candidate.parents and candidate != root:
        raise ValueError("Source-material path is outside the workspace folder.")
    if not candidate.is_file():
        raise ValueError(f"Source-material file not found: {relpath}")
    return candidate


def extract_folder_file(relpath: str) -> tuple[str, list[str]]:
    path = _resolve_folder_file(relpath)
    return extract_text_from_bytes(path.name, path.read_bytes())


def parse_source_files_json(value: str) -> list[str]:
    if not value:
        return []
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError:
        return []
    if not isinstance(parsed, list):
        return []
    return [str(x) for x in parsed if str(x).strip()]


def build_source_context(
    *,
    pasted_text: str = "",
    folder_files: list[str] | None = None,
    uploaded_files: list | None = None,
    strict: bool = False,
) -> dict:
    materials: list[dict] = []
    warnings: list[str] = []
    errors: list[str] = []

    pasted = (pasted_text or "").strip()
    if pasted:
        text, truncated = _truncate(_collapse_ws(pasted))
        if truncated:
            warnings.append(
                f"Pasted source material was truncated to {MAX_EXTRACTED_CHARS:,} characters."
            )
        materials.append({
            "title": "Pasted source material",
            "source": "pasted",
            "text": text,
            "chars": len(text),
            "tokens_est": estimate_text_tokens(text),
        })

    for relpath in folder_files or []:
        try:
            text, file_warnings = extract_folder_file(relpath)
            warnings.extend(file_warnings)
            materials.append({
                "title": Path(relpath).name,
                "source": f"Source Materials/{relpath}",
                "text": text,
                "chars": len(text),
                "tokens_est": estimate_text_tokens(text),
            })
        except Exception as e:
            errors.append(f"{relpath}: {e}")

    for upload in uploaded_files or []:
        filename = getattr(upload, "filename", "") or "uploaded file"
        if not filename:
            continue
        try:
            fileobj = getattr(upload, "file", None)
            if fileobj is None:
                continue
            data = fileobj.read()
            text, file_warnings = extract_text_from_bytes(filename, data)
            warnings.extend(file_warnings)
            materials.append({
                "title": Path(filename).name,
                "source": "uploaded file",
                "text": text,
                "chars": len(text),
                "tokens_est": estimate_text_tokens(text),
            })
        except Exception as e:
            errors.append(f"{filename}: {e}")

    if strict and errors:
        raise ValueError("; ".join(errors))

    total_text = "\n\n".join(m["text"] for m in materials)
    total_tokens = estimate_text_tokens(total_text)
    return {
        "materials": materials,
        "warnings": warnings,
        "errors": errors,
        "tokens_est": total_tokens,
        "chars": len(total_text),
    }


def context_warnings(context: dict) -> list[str]:
    warnings = list((context or {}).get("warnings") or [])
    tokens = int((context or {}).get("tokens_est") or 0)
    if tokens >= 100_000:
        warnings.append(
            "This looks book-sized. Whole books can get expensive quickly, especially "
            "if you retry or compare multiple models. Use excerpts when possible."
        )
    elif tokens >= 50_000:
        warnings.append(
            "This is a very large source context. Use an excerpt unless the whole text "
            "is truly needed for scoring."
        )
    elif tokens >= 10_000:
        warnings.append(
            "This is a long source context. Cost is still estimated as fresh input; "
            "provider caching is not guaranteed."
        )
    return warnings


def response_preset(kind: str) -> dict:
    return RESPONSE_PRESETS.get((kind or "scr").lower(), RESPONSE_PRESETS["scr"])


def synthetic_student_response(kind: str) -> str:
    preset = response_preset(kind)
    return ("student response " * int(preset["response_words"])).strip()
