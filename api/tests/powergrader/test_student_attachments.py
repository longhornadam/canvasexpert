import io
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from api.powergrader import new_quiz_fetch as nq
from api.powergrader import student_attachments as att


def _png_bytes(color="red"):
    out = io.BytesIO()
    Image.new("RGB", (3, 3), color).save(out, format="PNG")
    return out.getvalue()


def test_docx_preserves_body_order_and_inline_image_as_safe_derivative():
    doc = Document()
    paragraph = doc.add_paragraph("Before image ")
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(_png_bytes()))
    paragraph.add_run(" After image")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    raw = io.BytesIO()
    doc.save(raw)

    result = att.route_bytes("response.docx", raw.getvalue())
    assert result["extraction_status"] == "extracted"
    text = result["text"]
    assert text.index("Before image") < text.index("Inline image") < text.index("After image")
    assert "A | B" in text
    assert result["media_derivative"]


def test_raster_derivative_is_reencoded_without_source_metadata():
    result = att.route_bytes("photo.png", _png_bytes())
    assert result["ai_eligible"] is True
    assert result["media_type"] == "image/png"
    with Image.open(io.BytesIO(result["media_derivative"])) as image:
        assert not image.info


def test_unsupported_or_over_budget_evidence_is_held_without_truncation():
    result = att.route_bytes("essay.txt", b"x" * 12, max_ai_chars=5)
    assert result["extraction_status"] == "exceeds_ai_budget"
    assert len(result["text"]) == 12
    decision = att.eligibility_decision([result])
    assert decision["held"] is True
    assert "exceeds_ai_budget" in decision["reasons"][0]


class _DownloadResponse:
    status_code = 200

    def iter_content(self, chunk_size):
        yield b"synthetic bytes"


class _HttpRedirectResponse(_DownloadResponse):
    url = "http://storage.invalid/final"


class _CleanClient:
    def __init__(self):
        self.calls = []

    def get(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return _DownloadResponse()


def test_signed_download_is_atomic_and_uses_clean_client(tmp_path):
    client = _CleanClient()
    dest = tmp_path / "upload.bin"
    result = nq._download_signed_url(
        "https://storage.invalid/synthetic", str(dest),
        http_session=client, declared_size=len(b"synthetic bytes"), free_space=10_000,
    )
    assert dest.read_bytes() == b"synthetic bytes"
    assert result["actual_size"] == len(b"synthetic bytes")
    assert "headers" not in client.calls[0][1]
    assert not (tmp_path / "upload.bin.partial").exists()


def test_signed_download_size_mismatch_cleans_partial(tmp_path):
    client = _CleanClient()
    dest = tmp_path / "upload.bin"
    with pytest.raises(ValueError, match="size"):
        nq._download_signed_url(
            "https://storage.invalid/synthetic", str(dest),
            http_session=client, declared_size=999, free_space=10_000,
        )
    assert not dest.exists()
    assert not (tmp_path / "upload.bin.partial").exists()


def test_signed_download_rejects_non_https_redirect(tmp_path):
    class _RedirectClient:
        def get(self, url, **kwargs):
            return _HttpRedirectResponse()

    dest = tmp_path / "upload.bin"
    with pytest.raises(ValueError, match="HTTPS"):
        nq._download_signed_url(
            "https://storage.invalid/synthetic", str(dest),
            http_session=_RedirectClient(), free_space=10_000,
        )
    assert not dest.exists()
    assert not (tmp_path / "upload.bin.partial").exists()
