from docx import Document

from api.webui import source_materials, workspace
from api.webui.source_material_extractors import extract_text_from_bytes as _extract_text_from_bytes


def test_extract_plain_text_direct():
    """The extractor module can be imported directly for dependency-light cases."""
    text, warnings = _extract_text_from_bytes("hello.txt", b"Hello, world!")
    assert warnings == []
    assert text == "Hello, world!"


def test_extract_html_direct():
    """HTML stripping works when importing the extractor module directly."""
    text, warnings = _extract_text_from_bytes("page.html", b"<p>Hello<br>world</p>")
    assert warnings == []
    assert "Hello" in text
    assert "world" in text


def test_extract_docx_text_from_uploaded_bytes(tmp_path):
    path = tmp_path / "story.docx"
    doc = Document()
    doc.add_paragraph("A short story paragraph.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Character"
    table.cell(0, 1).text = "Motivation"
    doc.save(path)

    text, warnings = source_materials.extract_text_from_bytes(path.name, path.read_bytes())

    assert warnings == []
    assert "A short story paragraph." in text
    assert "Character | Motivation" in text


def test_build_source_context_from_paste_and_folder_file(tmp_path, monkeypatch):
    root = tmp_path / "CanvasExpert"
    source_dir = root / "Source Materials"
    source_dir.mkdir(parents=True)
    (source_dir / "passage.html").write_text(
        "<h1>Passage</h1><p>The river changed overnight.</p>",
        encoding="utf-8",
    )
    monkeypatch.setattr(workspace, "workspace_root", lambda: str(root))

    ctx = source_materials.build_source_context(
        pasted_text="Pasted excerpt.",
        folder_files=["passage.html"],
    )

    assert len(ctx["materials"]) == 2
    assert ctx["tokens_est"] > 0
    assert "The river changed overnight." in ctx["materials"][1]["text"]


def test_context_warnings_call_out_book_sized_text():
    ctx = {"tokens_est": 110_000, "warnings": []}

    warnings = source_materials.context_warnings(ctx)

    assert any("book-sized" in w for w in warnings)
