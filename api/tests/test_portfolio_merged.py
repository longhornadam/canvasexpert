"""Offline tests for the merged-portfolio core: entry building, chronological
sorting, and attachment embedding (image / docx / pdf-fallback / unknown).

No live Canvas, no PII (uses the synthetic NQ fixture + generated files). The live
fetch (portfolio_service.build_merged_portfolios) is thin glue over the proven
student_packet submission fetch and is validated against a real course manually.
"""
import os

from docx import Document
from PIL import Image

from api import portfolio
from api.nq_report import parse_student_analysis_file

FIXTURE = os.path.join(os.path.dirname(__file__), "fixtures",
                       "student_analysis_sample.csv")


def test_nq_entries_shape():
    data = parse_student_analysis_file(FIXTURE)
    ada = data["students"][0]
    entries = portfolio.nq_entries(data, ada, "THG")
    # 1 MC-summary entry + 2 constructed responses
    assert len(entries) == 3
    summary = entries[0]
    assert summary["title"] == "THG — multiple choice"
    assert "2 correct" in summary["response_text"]
    written = entries[1:]
    assert all(e["source"] == "New Quiz" for e in written)
    assert {e["possible"] for e in written} == {10.0, 8.0}   # inferred per-item maxima
    assert any("autumn" in e["response_text"] for e in written)


def test_sort_entries_chronological_undated_last():
    rows = [{"date": "2026-02-01", "source": "A"},
            {"date": "", "source": "B"},
            {"date": "2026-01-01", "source": "C"}]
    ordered = [e["date"] for e in portfolio.sort_entries(rows)]
    assert ordered == ["2026-01-01", "2026-02-01", ""]


def test_render_merged_with_attachments(tmp_path):
    # Synthetic attachments: a real PNG, a real DOCX, a fake PDF, an unknown type.
    png = tmp_path / "photo.png"
    Image.new("RGB", (4, 4), (120, 180, 240)).save(str(png))

    sub = Document()
    sub.add_paragraph("Handwritten transcription paragraph.")
    docx_path = tmp_path / "scan.docx"
    sub.save(str(docx_path))

    pdf_path = tmp_path / "essay.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 not-a-real-pdf")
    other = tmp_path / "notes.xyz"
    other.write_text("x")

    entries = [
        {"date": "2026-01-15", "source": "Assignment", "title": "Personal Narrative",
         "prompt": "Write about a turning point.", "response_text": "My story...",
         "attachments": [str(png), str(docx_path), str(pdf_path), str(other)],
         "earned": 18.0, "possible": 20.0},
    ]
    dest = tmp_path / "merged.docx"
    portfolio.render_merged_docx("Sample Student", entries, str(dest))
    assert dest.exists() and dest.stat().st_size > 0

    text = "\n".join(p.text for p in Document(str(dest)).paragraphs)
    assert "Sample Student" in text                       # title
    assert "Personal Narrative" in text                   # entry heading
    assert "Handwritten transcription paragraph." in text  # docx extracted
    assert "[Attached PDF: essay.pdf]" in text             # pdf fallback (pypdf absent / fake)
    assert "[Attached file: notes.xyz]" in text            # unknown reference
    assert "(20 possible)" in text                         # earned (possible)
