"""Student Reports — compile one student's work + standing into a local packet.

Folder layout, per student, per course:  <root>/<Student>/<Course>/{Assignments,Info}
Neutral, label-free output (see SPEC_Student_Reports Language rules). Local-only.
"""
import json
import os
from datetime import datetime

import requests
from docx import Document

from api.submission_transport import download_binary as _download_binary, get_all_pages as _get_all_pages
from api.webui.workspace import safe_component


def safe_name(value, max_len=80):
    return safe_component(value, max_len=max_len)


def _student_file_tag(user, fallback_user_id=None):
    user = user or {}
    raw = user.get("sortable_name") or user.get("name") or f"user_{fallback_user_id}"
    return safe_name(raw, 40)


def _work_filename(assignment_name, student_tag, ext="", detail=""):
    stem = safe_name(assignment_name, 110) + " - " + safe_name(student_tag, 40)
    return stem + (" - " + safe_name(detail, 40) if detail else "") + (ext or "")


def _reserve_filename(filename, used):
    stem, ext = os.path.splitext(filename)
    directory = used.get("__directory__", "") if isinstance(used, dict) else ""
    names = used.setdefault("__names__", set()) if isinstance(used, dict) else used
    candidate, number = filename, 2
    while candidate in names or (directory and os.path.exists(os.path.join(directory, candidate))):
        candidate = f"{stem} ({number}){ext}"; number += 1
    names.add(candidate)
    return candidate

# Section keys the UI offers (order preserved in the document):
SECTIONS = ["standing", "late", "adjustments", "comments", "work"]

DOWNLOADABLE = {"online_text_entry", "online_upload", "online_url"}


# ── DOCX ─────────────────────────────────────────────────────────────────


def _render_info_docx(dest, student_name, course_name, blocks):
    """blocks: ordered list of (heading, kind, payload):
       kind 'table' → payload {cols:[...], rows:[[...]]}; kind 'lines' → [str]."""
    doc = Document()
    doc.add_heading(student_name, level=0)
    doc.add_paragraph(f"{course_name} · generated {datetime.now():%b %d, %Y}")
    for heading, kind, payload in blocks:
        doc.add_heading(heading, level=2)
        if kind == "table" and payload and payload["rows"]:
            cols = payload["cols"]
            t = doc.add_table(rows=1, cols=len(cols))
            t.style = "Light Grid Accent 1"
            for i, c in enumerate(cols):
                t.rows[0].cells[i].text = str(c)
            for row in payload["rows"]:
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = "" if val is None else str(val)
        elif kind == "lines" and payload:
            for ln in payload:
                doc.add_paragraph(ln, style="List Bullet")
        else:
            doc.add_paragraph("None on record.")
    doc.save(dest)


# ── Data shaping (neutral language lives HERE) ──────────────────────────────


def _days_late(seconds):
    try:
        return max(1, round(int(seconds) / 86400))
    except (TypeError, ValueError):
        return None


def _info_blocks(subs, curve_rows, sections):
    standing, late, adj, comments = [], [], [], []
    for s in subs:
        a = s.get("assignment") or {}
        name = a.get("name", "?")
        pts = a.get("points_possible", "")
        state = "Excused" if s.get("excused") else (s.get("workflow_state") or "")
        standing.append([name, str(pts),
                         str(s.get("score", "") if s.get("score") is not None else ""),
                         state, (s.get("submitted_at") or "")[:10]])
        # late / extended due date — factual, no "accommodation"
        if s.get("excused"):
            late.append(f'{name}: Excused')
        else:
            base_due = (a.get("due_at") or "")[:10]
            cached_due = (s.get("cached_due_date") or "")[:10]
            if cached_due and base_due and cached_due > base_due:
                late.append(f'{name}: Due date extended to {cached_due}')
            if s.get("late") and s.get("seconds_late"):
                d = _days_late(s["seconds_late"])
                if d:
                    late.append(f'{name}: Submitted {d} day(s) late')
        for c in (s.get("submission_comments") or []):
            who = (c.get("author_name") or "").strip()
            when = (c.get("created_at") or "")[:10]
            comments.append(f'{when} — {who}: {c.get("comment", "")}')
    for cr in curve_rows:
        adj.append(cr)   # pre-formatted neutral strings, built in build_packet
    blocks = []
    if "standing" in sections:
        blocks.append(("Standing", "table",
                       {"cols": ["Assignment", "Pts Available", "Score", "Status", "Submitted"],
                        "rows": standing}))
    if "late" in sections:
        blocks.append(("Late & extended due dates", "lines", late))
    if "adjustments" in sections:
        blocks.append(("Adjustments", "lines", adj))
    if "comments" in sections:
        blocks.append(("Comments", "lines", comments))
    return blocks


def _signature(subs, curve_rows):
    """Cheap change-detector for dedupe: latest submission marker + counts."""
    last = max([(s.get("submitted_at") or "") for s in subs] + [""])
    graded = sum(1 for s in subs if s.get("workflow_state") == "graded")
    return f"{last}|{graded}|{len(curve_rows)}"


# ── Per-student build ───────────────────────────────────────────────────────


def build_packet(user_id, student_name, sections, courses, base, token,
                 reports_root, curve_events, skip_unchanged=False):
    """Generator of progress strings. `courses` = [{id, name}] to consider.
    Final line: 'FOLDER: <student root>'. Set skip_unchanged for the routine path."""
    session = requests.Session()
    session.headers["Authorization"] = f"Bearer {token}"
    stu_root = os.path.join(reports_root, safe_name(student_name))
    file_tag = _student_file_tag({"name": student_name}, user_id)
    os.makedirs(stu_root, exist_ok=True)
    man_path = os.path.join(stu_root, "_manifest.json")
    try:
        with open(man_path, encoding="utf-8") as f:
            manifest = json.load(f)
    except Exception:
        manifest = {}

    yield f"Student: {student_name}"
    any_course = False
    for c in courses:
        cid, cname = str(c["id"]), c["name"]
        try:
            subs = _get_all_pages(
                session, f"{base}/api/v1/courses/{cid}/students/submissions",
                {"student_ids[]": str(user_id),
                 "include[]": ["assignment", "submission_comments"],
                 "per_page": 100})
        except requests.HTTPError as e:
            code = e.response.status_code if e.response is not None else "?"
            yield f"· {cname}: no access (HTTP {code}) — skipped"
            continue                         # e.g. 403 where the teacher can't read submissions
        except Exception as e:
            yield f"· {cname}: skipped ({e})"
            continue
        subs = [s for s in (subs or []) if (s.get("assignment") or {}).get("id")]
        if not subs:
            continue                         # student not in this course
        any_course = True
        # curve adjustments for this student+course (local records, neutral phrasing)
        curve_rows = []
        for ev in curve_events:
            if str(ev.get("course_id")) != cid or ev.get("reverted"):
                continue
            for st in (ev.get("students") or []):
                if str(st.get("user_id")) == str(user_id):
                    old, new = st.get("original_score"), st.get("curved_score")
                    when = (ev.get("applied_at") or "")[:10]
                    curve_rows.append(
                        f'{ev.get("assignment_name", ev.get("assignment_id"))}: '
                        f'Score adjusted via curve on {when}: {old} → {new}')

        sig = _signature(subs, curve_rows)
        if skip_unchanged and manifest.get(cid, {}).get("signature") == sig:
            yield f"· {cname}: no change since last packet — skipped"
            continue

        course_dir = os.path.join(stu_root, safe_name(cname))
        info_dir = os.path.join(course_dir, "Info")
        os.makedirs(info_dir, exist_ok=True)

        # work samples (original formats) → Assignments/
        if "work" in sections:
            asg_dir = os.path.join(course_dir, "Assignments")
            os.makedirs(asg_dir, exist_ok=True)
            used_filenames = set()
            n = 0
            for s in subs:
                a = s.get("assignment") or {}
                assignment_name = a.get("name", "work")
                body = (s.get("body") or "").strip()
                if body:
                    fname = _reserve_filename(
                        _work_filename(assignment_name, file_tag, ".html"),
                        used_filenames,
                    )
                    with open(os.path.join(asg_dir, fname), "w", encoding="utf-8") as f:
                        f.write(body)
                    n += 1
                url = (s.get("url") or "").strip()
                if url:
                    fname = _reserve_filename(
                        _work_filename(assignment_name, file_tag, ".txt", "URL"),
                        used_filenames,
                    )
                    with open(os.path.join(asg_dir, fname), "w", encoding="utf-8") as f:
                        f.write(url + "\n")
                    n += 1
                for att in (s.get("attachments") or []):
                    orig = att.get("filename") or att.get("display_name") or "file"
                    detail, ext = os.path.splitext(orig)
                    fname = _reserve_filename(
                        _work_filename(assignment_name, file_tag, ext, detail or "file"),
                        used_filenames,
                    )
                    dest = os.path.join(asg_dir, fname)
                    try:
                        _download_binary(session, att["url"], dest)
                        n += 1
                    except Exception as e:
                        yield f"  !! {cname}/{safe_name(assignment_name)}: {e}"
            yield f"✓ {cname}: {n} work file(s)"

        # Info DOCX
        blocks = _info_blocks(subs, curve_rows, sections)
        out = os.path.join(info_dir, f"{safe_name(student_name)} - {safe_name(cname)} "
                                     f"- {datetime.now():%Y-%m-%d}.docx")
        _render_info_docx(out, student_name, cname, blocks)
        yield f"✓ {cname}: Info document written"
        manifest[cid] = {"signature": sig,
                         "last_run": datetime.now().isoformat(timespec="seconds")}

    if not any_course:
        yield "· student not found in any selected course"
    with open(man_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    yield f"FOLDER: {stu_root}"
